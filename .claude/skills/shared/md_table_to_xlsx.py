#!/usr/bin/env python3
"""
md_table_to_xlsx.py — Convert markdown table(s) to an XLSX workbook.

Usage:
    python md_table_to_xlsx.py <input.md> <output.xlsx>

Input format:
    One or more standard pipe-delimited markdown tables. Separate tables intended
    for different sheets with a "## Sheet: <name>" heading immediately before each table.
    If no heading is present, the single table is written to a sheet named "Sheet1".

    Example (single sheet):
        | ID | Name | Status |
        |----|------|--------|
        | 1  | Foo  | Pass   |

    Example (multiple sheets):
        ## Sheet: Test Cases
        | ID | Name |
        |----|------|
        | 1  | Foo  |

        ## Sheet: Summary
        | Category | Count |
        |----------|-------|
        | Total    | 1     |

Output:
    Prints {"status": "success", "file": "<output.xlsx>"} on success,
    or {"status": "error", "message": "..."} on failure, then exits 1.
"""

import sys
import json
import re
import os

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print(json.dumps({
        "status": "error",
        "message": "openpyxl not installed. Run: pip install openpyxl"
    }))
    sys.exit(1)


# ── Style constants ────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
DATA_FONT   = Font(size=10, name="Calibri")
WRAP_ALIGN  = Alignment(wrap_text=True, vertical="top")
TOP_ALIGN   = Alignment(vertical="top", wrap_text=False)

ALT_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

MAX_COL_WIDTH = 60
MIN_COL_WIDTH = 10
DATA_ROW_HEIGHT = 55


# ── Parsing ────────────────────────────────────────────────────────────────────

def _split_row(line: str) -> list[str]:
    """Split a pipe-delimited markdown row into a list of stripped cell strings."""
    return [cell.strip().replace("\\n", "\n") for cell in line.strip().strip("|").split("|")]


def _is_separator(line: str) -> bool:
    """Return True if the line is a markdown table separator (e.g. |---|---|)."""
    stripped = line.strip().strip("|")
    return bool(re.fullmatch(r"[\s\-:|]+", stripped)) and "-" in stripped


def _parse_table_block(lines: list[str]) -> tuple[list[str], list[list[str]]] | None:
    """
    Parse a single markdown table from a list of lines.
    Returns (headers, data_rows) or None if no valid table is found.
    """
    table_lines = [l for l in lines if l.strip().startswith("|")]
    if len(table_lines) < 2:
        return None

    headers = _split_row(table_lines[0])
    data_rows = []
    for line in table_lines[1:]:
        if _is_separator(line):
            continue
        row = _split_row(line)
        # Pad or trim row to match header column count
        while len(row) < len(headers):
            row.append("")
        data_rows.append(row[: len(headers)])

    return headers, data_rows


def parse_md_tables(text: str) -> list[tuple[str, list[str], list[list[str]]]]:
    """
    Parse all markdown tables from text, respecting "## Sheet: <name>" headings.
    Returns a list of (sheet_name, headers, data_rows).
    """
    sheets: list[tuple[str, list[str], list[list[str]]]] = []
    current_name = "Sheet1"
    current_lines: list[str] = []

    for line in text.splitlines():
        heading_match = re.match(r"^##\s+Sheet:\s*(.+)", line.strip())
        if heading_match:
            parsed = _parse_table_block(current_lines)
            if parsed:
                safe_name = _safe_sheet_name(current_name, [s[0] for s in sheets])
                sheets.append((safe_name, parsed[0], parsed[1]))
            current_name = heading_match.group(1).strip()
            current_lines = []
        else:
            current_lines.append(line)

    # Flush the last block
    parsed = _parse_table_block(current_lines)
    if parsed:
        safe_name = _safe_sheet_name(current_name, [s[0] for s in sheets])
        sheets.append((safe_name, parsed[0], parsed[1]))

    return sheets


def _safe_sheet_name(name: str, existing: list[str]) -> str:
    """Truncate and deduplicate Excel sheet names (max 31 chars, no special chars)."""
    clean = re.sub(r"[\\/*?:\[\]]", "-", name)[:31]
    if clean not in existing:
        return clean
    # Append a suffix to deduplicate
    for i in range(2, 100):
        candidate = f"{clean[: 28]}-{i}"
        if candidate not in existing:
            return candidate
    return clean


# ── XLSX writing ──────────────────────────────────────────────────────────────

def _col_width(header: str, max_data_len: int) -> float:
    """Compute a sensible column width clamped to [MIN, MAX]."""
    ideal = max(len(header) + 2, min(max_data_len + 2, MAX_COL_WIDTH))
    return float(max(MIN_COL_WIDTH, min(ideal, MAX_COL_WIDTH)))


def write_xlsx(
    sheets: list[tuple[str, list[str], list[list[str]]]],
    output_path: str,
) -> None:
    """Write parsed sheet data to an XLSX workbook at output_path."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove the default blank sheet

    for sheet_name, headers, rows in sheets:
        ws = wb.create_sheet(title=sheet_name)

        # Header row
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = WRAP_ALIGN

        # Data rows
        for row_idx, row in enumerate(rows, start=2):
            fill = ALT_FILL if row_idx % 2 == 0 else None
            for col_idx in range(1, len(headers) + 1):
                value = row[col_idx - 1] if col_idx - 1 < len(row) else ""
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = DATA_FONT
                cell.alignment = WRAP_ALIGN
                if fill:
                    cell.fill = fill

        # Freeze header row
        ws.freeze_panes = "A2"

        # Column widths — use longest single line within each cell, not total length
        for col_idx, header in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            max_data = max(
                (
                    max(len(ln) for ln in str(row[col_idx - 1]).split("\n"))
                    if col_idx - 1 < len(row)
                    else 0
                    for row in rows
                ),
                default=0,
            )
            ws.column_dimensions[col_letter].width = _col_width(header, max_data)

        # Row heights — scale with the number of wrapped lines in the tallest cell
        ws.row_dimensions[1].height = 20
        for row_idx, row in enumerate(rows, start=2):
            max_lines = max(
                (str(row[ci]).count("\n") + 1 if ci < len(row) else 1
                 for ci in range(len(headers))),
                default=1,
            )
            ws.row_dimensions[row_idx].height = max(DATA_ROW_HEIGHT, max_lines * 15 + 5)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)


# ── Entry point ───────────────────────────────────────────────────────────────

def main() -> None:
    if len(sys.argv) == 2 and sys.argv[1] in ("--help", "-h"):
        print(__doc__)
        sys.exit(0)

    if len(sys.argv) != 3:
        print(json.dumps({
            "status": "error",
            "message": (
                f"Usage: python {os.path.basename(sys.argv[0])} <input.md> <output.xlsx>"
            ),
        }))
        sys.exit(1)

    input_path, output_path = sys.argv[1], sys.argv[2]

    try:
        with open(input_path, "r", encoding="utf-8") as fh:
            text = fh.read()
    except OSError as exc:
        print(json.dumps({"status": "error", "message": f"Cannot read input file: {exc}"}))
        sys.exit(1)

    sheets = parse_md_tables(text)
    if not sheets:
        print(json.dumps({
            "status": "error",
            "message": "No markdown tables found in the input file.",
        }))
        sys.exit(1)

    try:
        write_xlsx(sheets, output_path)
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"status": "error", "message": f"Failed to write XLSX: {exc}"}))
        sys.exit(1)

    print(json.dumps({"status": "success", "file": output_path}))


if __name__ == "__main__":
    main()
