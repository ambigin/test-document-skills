"""
Converter tests for skills/*/scripts/md_table_to_xlsx.py.
Run with: pytest tests/test_converter.py

Each test is parameterised over all seven skill converters so any
divergence between copies is caught immediately.
"""

import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).parent.parent
CONVERTERS = sorted(REPO.glob("skills/*/scripts/md_table_to_xlsx.py"))


def run(converter: Path, md: str, tmp_path: Path):
    md_file = tmp_path / "input.md"
    xlsx_file = tmp_path / "output.xlsx"
    md_file.write_text(md, encoding="utf-8")
    r = subprocess.run(
        [sys.executable, str(converter), str(md_file), str(xlsx_file)],
        capture_output=True, text=True,
    )
    try:
        return json.loads(r.stdout), xlsx_file
    except json.JSONDecodeError:
        return {"status": "error", "raw": r.stdout + r.stderr}, xlsx_file


@pytest.fixture(
    params=[str(c) for c in CONVERTERS],
    ids=[c.parent.parent.name for c in CONVERTERS],
)
def converter(request):
    return Path(request.param)


# ── Basic output ───────────────────────────────────────────────────────────────

def test_single_table_success(converter, tmp_path):
    r, xlsx = run(converter, "| A | B |\n|---|---|\n| 1 | 2 |\n", tmp_path)
    assert r["status"] == "success" and xlsx.exists()


def test_default_sheet_name(converter, tmp_path):
    _, xlsx = run(converter, "| A | B |\n|---|---|\n| 1 | 2 |\n", tmp_path)
    assert openpyxl.load_workbook(xlsx).sheetnames == ["Sheet1"]


def test_multi_sheet_names(converter, tmp_path):
    md = "## Sheet: Alpha\n| X |\n|---|\n| v |\n\n## Sheet: Beta\n| Y |\n|---|\n| w |\n"
    _, xlsx = run(converter, md, tmp_path)
    assert openpyxl.load_workbook(xlsx).sheetnames == ["Alpha", "Beta"]


def test_header_values(converter, tmp_path):
    _, xlsx = run(converter, "| ID | Name | Status |\n|---|---|---|\n| 1 | Foo | Pass |\n", tmp_path)
    ws = openpyxl.load_workbook(xlsx).active
    assert [ws.cell(1, c).value for c in range(1, 4)] == ["ID", "Name", "Status"]


def test_data_row_values(converter, tmp_path):
    _, xlsx = run(converter, "| ID | Name |\n|---|---|\n| 42 | Bar |\n", tmp_path)
    ws = openpyxl.load_workbook(xlsx).active
    assert str(ws.cell(2, 1).value) == "42"
    assert ws.cell(2, 2).value == "Bar"


def test_newline_escape_unescaped(converter, tmp_path):
    _, xlsx = run(converter, "| Steps |\n|---|\n| Step 1\\nStep 2 |\n", tmp_path)
    val = openpyxl.load_workbook(xlsx).active.cell(2, 1).value
    assert "\n" in val, f"Expected real newline in cell, got: {repr(val)}"


# ── Smoke test with realistic input ───────────────────────────────────────────

SAMPLE_MD = """\
## Sheet: Test Cases

| Test Case ID | Test Case Name | Test Steps | Expected Result | Status |
|---|---|---|---|---|
| TC-001 | Happy path login | 1. Open /login\\n2. Enter valid email\\n3. Enter valid password\\n4. Click Submit | User is redirected to dashboard | Not Executed |
| TC-002 | Login with wrong password | 1. Open /login\\n2. Enter valid email\\n3. Enter wrong password\\n4. Click Submit | Error message "Invalid credentials" shown | Not Executed |

## Sheet: Summary

| Category | Count |
|---|---|
| Total Test Cases | 2 |
| Positive Flow | 1 |
| Negative Flow | 1 |
"""


def test_sample_produces_correct_sheets(converter, tmp_path):
    r, xlsx = run(converter, SAMPLE_MD, tmp_path)
    assert r["status"] == "success"
    wb = openpyxl.load_workbook(xlsx)
    assert "Test Cases" in wb.sheetnames
    assert "Summary" in wb.sheetnames


def test_sample_row_count(converter, tmp_path):
    _, xlsx = run(converter, SAMPLE_MD, tmp_path)
    ws = openpyxl.load_workbook(xlsx)["Test Cases"]
    assert ws.max_row == 3  # 1 header + 2 data rows


# ── Copy consistency ───────────────────────────────────────────────────────────

def test_all_converter_copies_identical():
    contents = [c.read_text() for c in CONVERTERS]
    assert len(set(contents)) == 1, (
        "Converter copies differ between skills — update all seven after any change: "
        + str([c.parent.parent.name for c in CONVERTERS])
    )
