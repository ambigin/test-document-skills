#!/usr/bin/env python3
r"""
md_table_to_xlsx.py — Convert markdown table(s) to an XLSX workbook.

Usage:
    python md_table_to_xlsx.py <input.md> <output.xlsx>

Input format:
    One or more standard pipe-delimited markdown tables. Separate tables intended
    for different sheets with a "## Sheet: <name>" heading immediately before each table.
    If no heading is present, the single table is written to a sheet named "Sheet1".
    Put exactly one table under each heading.

    Every table row must be on a single line. Inside a cell, write a line break
    as \n and a literal pipe as \|.

    Example (single sheet):
        | ID | Name | Status |
        |----|------|--------|
        | 1  | Foo  | Pass   |

    Example (multiple sheets):
        ## Sheet: Test Cases
        | ID | Name |
        |----|------|
        | 1  | Foo  |

        ## Sheet: Summary
        | Category | Count |
        |----------|-------|
        | Total    | 1     |

Output:
    Prints {"status": "success", "file": "<output.xlsx>"} on success, plus a
    "warnings" list if any content was dropped (extra cells, a second table under
    one heading, or lines inside a table that don't start with "|"),
    or {"status": "error", "message": "..."} on failure, then exits 1.

    Cells that start with "=" are written as plain text, never as formulas.
"""

from __future__ import annotations

import sys
import json
import math
import re
import os

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print(json.dumps({
        "status": "error",
        "message": "openpyxl not installed. Run: pip install openpyxl"
    }))
    sys.exit(1)


# ── Style constants ────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
DATA_FONT   = Font(size=10, name="Calibri")
WRAP_ALIGN  = Alignment(wrap_text=True, vertical="top")
TOP_ALIGN   = Alignment(vertical="top", wrap_text=False)

ALT_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

MAX_COL_WIDTH = 60
MIN_COL_WIDTH = 10
DATA_ROW_HEIGHT = 55
MAX_ROW_HEIGHT = 409  # Excel's maximum row height, in points
LINE_HEIGHT = 15


# ── Parsing ────────────────────────────────────────────────────────────────────

# A cell boundary is a "|" that is not escaped as "\|".
_UNESCAPED_PIPE = re.compile(r"(?<!\\)\|")


def _split_row(line: str) -> list[str]:
    r"""Split a markdown row on unescaped pipes; turn \| into | and \n into a newline."""
    row = line.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|") and not row.endswith("\\|"):
        row = row[:-1]
    return [
        cell.strip().replace("\\|", "|").replace("\\n", "\n")
        for cell in _UNESCAPED_PIPE.split(row)
    ]


def _is_separator(line: str) -> bool:
    """Return True if the line is a markdown table separator (e.g. |---|---|)."""
    stripped = line.strip().strip("|")
    return bool(re.fullmatch(r"[\s\-:|]+", stripped)) and "-" in stripped


def _parse_table_block(
    lines: list[str], sheet_name: str, warnings: list[str]
) -> tuple[list[str], list[list[str]]] | None:
    """
    Parse a single markdown table from a list of lines.
    Returns (headers, data_rows) or None if no valid table is found.
    Appends a message to `warnings` whenever content is dropped.
    """
    table_idx = [i for i, l in enumerate(lines) if l.strip().startswith("|")]
    if len(table_idx) < 2:
        return None

    # Text lines between table rows usually mean a cell contained a real line break
    stray = [
        l for l in lines[table_idx[0]:table_idx[-1]]
        if l.strip() and not l.strip().startswith("|")
    ]
    if stray:
        warnings.append(
            f"Sheet '{sheet_name}': {len(stray)} line(s) inside the table don't start "
            "with '|' and were ignored (write line breaks inside a cell as \\n)."
        )

    table_lines = [lines[i] for i in table_idx]
    headers = _split_row(table_lines[0])
    data_rows = []
    for pos, line in enumerate(table_lines[1:], start=1):
        if _is_separator(line):
            if pos > 1:
                warnings.append(
                    f"Sheet '{sheet_name}': a second table was merged in as data rows; "
                    "put a '## Sheet:' heading before each table."
                )
            continue
        row = _split_row(line)
        if len(row) > len(headers):
            warnings.append(
                f"Sheet '{sheet_name}': row '{row[0][:40]}' has {len(row)} cells but the "
                f"header has {len(headers)}; the extra cells were dropped "
                "(escape literal pipes as \\|)."
            )
        # Pad or trim row to match header column count
        while len(row) < len(headers):
            row.append("")
        data_rows.append(row[: len(headers)])

    return headers, data_rows


def parse_md_tables(
    text: str, warnings: list[str] | None = None
) -> list[tuple[str, list[str], list[list[str]]]]:
    """
    Parse all markdown tables from text, respecting "## Sheet: <name>" headings.
    Returns a list of (sheet_name, headers, data_rows). Messages about dropped
    content are appended to `warnings` when a list is passed in.
    """
    if warnings is None:
        warnings = []
    sheets: list[tuple[str, list[str], list[list[str]]]] = []
    current_name = "Sheet1"
    current_lines: list[str] = []

    for line in text.splitlines():
        heading_match = re.match(r"^##\s+Sheet:\s*(.+)", line.strip())
        if heading_match:
            parsed = _parse_table_block(current_lines, current_name, warnings)
            if parsed:
                safe_name = _safe_sheet_name(current_name, [s[0] for s in sheets])
                sheets.append((safe_name, parsed[0], parsed[1]))
            current_name = heading_match.group(1).strip()
            current_lines = []
        else:
            current_lines.append(line)

    # Flush the last block
    parsed = _parse_table_block(current_lines, current_name, warnings)
    if parsed:
        safe_name = _safe_sheet_name(current_name, [s[0] for s in sheets])
        sheets.append((safe_name, parsed[0], parsed[1]))

    return sheets


def _safe_sheet_name(name: str, existing: list[str]) -> str:
    """Truncate and deduplicate Excel sheet names (max 31 chars, no special chars)."""
    clean = re.sub(r"[\\/*?:\[\]]", "-", name)[:31]
    if clean not in existing:
        return clean
    # Append a suffix to deduplicate
    for i in range(2, 100):
        candidate = f"{clean[: 28]}-{i}"
        if candidate not in existing:
            return candidate
    return clean


# ── XLSX writing ──────────────────────────────────────────────────────────────

def _col_width(header: str, max_data_len: int) -> float:
    """Compute a sensible column width clamped to [MIN, MAX]."""
    ideal = max(len(header) + 2, min(max_data_len + 2, MAX_COL_WIDTH))
    return float(max(MIN_COL_WIDTH, min(ideal, MAX_COL_WIDTH)))


def _wrapped_lines(text: str, width: float) -> int:
    """Estimate how many lines a wrapped cell needs at the given column width."""
    chars_per_line = max(1, int(width) - 2)
    return sum(max(1, math.ceil(len(line) / chars_per_line)) for line in text.split("\n"))


def write_xlsx(
    sheets: list[tuple[str, list[str], list[list[str]]]],
    output_path: str,
) -> None:
    """Write parsed sheet data to an XLSX workbook at output_path."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove the default blank sheet

    for sheet_name, headers, rows in sheets:
        ws = wb.create_sheet(title=sheet_name)

        # Header row
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = WRAP_ALIGN

        # Data rows
        for row_idx, row in enumerate(rows, start=2):
            fill = ALT_FILL if row_idx % 2 == 0 else None
            for col_idx in range(1, len(headers) + 1):
                value = row[col_idx - 1] if col_idx - 1 < len(row) else ""
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if value.startswith("="):
                    cell.data_type = "s"  # keep text such as "=HYPERLINK(...)" from becoming a live formula
                cell.font = DATA_FONT
                cell.alignment = WRAP_ALIGN
                if fill:
                    cell.fill = fill

        # Freeze header row
        ws.freeze_panes = "A2"

        # Column widths — use longest single line within each cell, not total length
        widths: list[float] = []
        for col_idx, header in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            max_data = max(
                (
                    max(len(ln) for ln in str(row[col_idx - 1]).split("\n"))
                    if col_idx - 1 < len(row)
                    else 0
                    for row in rows
                ),
                default=0,
            )
            widths.append(_col_width(header, max_data))
            ws.column_dimensions[col_letter].width = widths[-1]

        # Row heights — scale with the wrapped lines in the tallest cell, capped at Excel's limit
        ws.row_dimensions[1].height = 20
        for row_idx, row in enumerate(rows, start=2):
            max_lines = max(
                (_wrapped_lines(str(row[ci]), widths[ci]) if ci < len(row) else 1
                 for ci in range(len(headers))),
                default=1,
            )
            ws.row_dimensions[row_idx].height = min(
                MAX_ROW_HEIGHT, max(DATA_ROW_HEIGHT, max_lines * LINE_HEIGHT + 5)
            )

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)


# ── Entry point ───────────────────────────────────────────────────────────────

def main() -> None:
    if len(sys.argv) == 2 and sys.argv[1] in ("--help", "-h"):
        print(__doc__)
        sys.exit(0)

    if len(sys.argv) != 3:
        print(json.dumps({
            "status": "error",
            "message": (
                f"Usage: python {os.path.basename(sys.argv[0])} <input.md> <output.xlsx>"
            ),
        }))
        sys.exit(1)

    input_path, output_path = sys.argv[1], sys.argv[2]

    try:
        with open(input_path, "r", encoding="utf-8") as fh:
            text = fh.read()
    except OSError as exc:
        print(json.dumps({"status": "error", "message": f"Cannot read input file: {exc}"}))
        sys.exit(1)

    warnings: list[str] = []
    sheets = parse_md_tables(text, warnings)
    if not sheets:
        print(json.dumps({
            "status": "error",
            "message": "No markdown tables found in the input file.",
        }))
        sys.exit(1)

    try:
        write_xlsx(sheets, output_path)
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"status": "error", "message": f"Failed to write XLSX: {exc}"}))
        sys.exit(1)

    result = {"status": "success", "file": output_path}
    if warnings:
        result["warnings"] = warnings
    print(json.dumps(result))


if __name__ == "__main__":
    main()
